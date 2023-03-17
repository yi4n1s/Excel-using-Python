from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
from openpyxl import load_workbook

wb = load_workbook('Pie.xlsx')
ws = wb.active

tab = Table(displayName='Table1', ref='A1:B5')
style = TableStyleInfo(name='TableStyleMedium9', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=True)

tab.TableStyleInfo = style
ws.add_table(tab)
#wb.save('table.xlsx')

img = Image('madecraft.jpg')
img.height = img.height * .25
img.width = img.width * .25
ws.add_image(img, 'C1')
wb.save('image.xlsx')